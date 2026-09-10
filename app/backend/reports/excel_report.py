"""Professional Excel exporter with a real Dashboard worksheet."""

from io import BytesIO
import re

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


DARK = "172033"
BLUE = "2563EB"
LIGHT = "EEF4FF"
BORDER = "D7DEE8"


def _safe_sheet(name):
    name = re.sub(r"[\\/*?:\[\]]", "_", str(name))
    return name[:31] or "Sheet"


def _numeric_columns(df):
    result = []
    for col in df.columns:
        s = df[col]
        if pd.api.types.is_numeric_dtype(s):
            result.append(col)
            continue
        converted = pd.to_numeric(
            s.astype(str).str.replace(",", "", regex=False)
            .str.replace("₹", "", regex=False)
            .str.replace("$", "", regex=False)
            .str.replace("%", "", regex=False),
            errors="coerce",
        )
        non_empty = s.notna().sum()
        if non_empty and converted.notna().sum() / non_empty >= 0.80:
            result.append(col)
    return result


def _style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=DARK)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 0
        for cell in ws[letter][:200]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 42)


def _write_df(writer, df, sheet):
    if not isinstance(df, pd.DataFrame):
        df = pd.DataFrame(df if df is not None else {})
    df.to_excel(writer, sheet_name=_safe_sheet(sheet), index=False)


def _make_dashboard(ws, df, kpis, quality, insights):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A6"
    ws.merge_cells("A1:H1")
    ws["A1"] = "CSV Auto-Analyzer Dashboard"
    ws["A1"].font = Font(size=20, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=DARK)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:H2")
    ws["A2"] = "Executive view generated automatically from the uploaded CSV"
    ws["A2"].font = Font(italic=True, color="64748B")

    # Dataset cards.
    cards = [
        ("Rows", f"{len(df):,}"),
        ("Columns", f"{len(df.columns):,}"),
        ("Numeric Fields", f"{len(_numeric_columns(df)):,}"),
        ("Health Score", f"{quality.get('health_score', '—')}/100"),
    ]
    for idx, (label, value) in enumerate(cards):
        col = 1 + idx * 2
        ws.cell(4, col, label)
        ws.cell(5, col, value)
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
        ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 1)
        for r in (4, 5):
            cell = ws.cell(r, col)
            cell.fill = PatternFill("solid", fgColor=LIGHT if r == 4 else "FFFFFF")
            cell.border = Border(*(Side(style="thin", color=BORDER),) * 4)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(4, col).font = Font(bold=True, color=DARK)
        ws.cell(5, col).font = Font(size=15, bold=True, color=BLUE)

    # KPI block.
    ws["A7"] = "Key Performance Indicators"
    ws["A7"].font = Font(size=13, bold=True, color=DARK)
    headers = ["KPI", "Value", "Description"]
    for c, h in enumerate(headers, 1):
        ws.cell(8, c, h)
    for c in range(1, 4):
        ws.cell(8, c).font = Font(bold=True, color="FFFFFF")
        ws.cell(8, c).fill = PatternFill("solid", fgColor=DARK)
    for r, item in enumerate((kpis or [])[:4], 9):
        ws.cell(r, 1, item.get("name", ""))
        ws.cell(r, 2, item.get("display_value", item.get("value", "")))
        ws.cell(r, 3, item.get("description", ""))

    # Insight block.
    start = 15
    ws.cell(start, 1, "Key Insights")
    ws.cell(start, 1).font = Font(size=13, bold=True, color=DARK)
    for i, insight in enumerate((insights or [])[:6], start + 1):
        ws.cell(i, 1, f"• {str(insight)}")
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)
        ws.cell(i, 1).alignment = Alignment(wrap_text=True, vertical="top")

    # Helper data for charts placed far right, then hidden.
    nums = _numeric_columns(df)
    cats = [c for c in df.columns if c not in nums]
    helper_col = 10
    chart_specs = []

    if cats and nums:
        cat, num = cats[0], nums[0]
        temp = df[[cat, num]].copy()
        temp[num] = pd.to_numeric(temp[num], errors="coerce")
        temp = temp.dropna().groupby(cat, as_index=False)[num].sum().sort_values(num, ascending=False).head(10)
        ws.cell(1, helper_col, cat); ws.cell(1, helper_col + 1, num)
        for r, row in enumerate(temp.itertuples(index=False), 2):
            ws.cell(r, helper_col, str(row[0])); ws.cell(r, helper_col + 1, float(row[1]))
        chart_specs.append(("bar", helper_col, len(temp) + 1, f"{num} by {cat}"))

    if nums:
        num = nums[0]
        temp = pd.to_numeric(df[num], errors="coerce").dropna().reset_index(drop=True)
        # Histogram-like buckets as an Excel bar chart.
        if len(temp) >= 2:
            bins = min(10, max(4, int(len(temp) ** 0.5)))
            counts = pd.cut(temp, bins=bins, include_lowest=True, labels=False)
            freq = counts.value_counts().sort_index()
            col = helper_col + 3
            ws.cell(1, col, "Range"); ws.cell(1, col + 1, "Count")
            for r, (bucket, count) in enumerate(freq.items(), 2):
                ws.cell(r, col, str(bucket + 1)); ws.cell(r, col + 1, int(count))
            chart_specs.append(("bar", col, len(freq) + 1, f"Distribution of {num}"))

    if len(nums) >= 2:
        a, b = nums[0], nums[1]
        temp = df[[a, b]].copy()
        temp[a] = pd.to_numeric(temp[a], errors="coerce")
        temp[b] = pd.to_numeric(temp[b], errors="coerce")
        temp = temp.dropna().head(30)
        col = helper_col + 6
        ws.cell(1, col, a); ws.cell(1, col + 1, b)
        for r, row in enumerate(temp.itertuples(index=False), 2):
            ws.cell(r, col, float(row[0])); ws.cell(r, col + 1, float(row[1]))
        chart_specs.append(("line", col, len(temp) + 1, f"{b} vs {a}"))

    # Add up to three actual Excel charts to the dashboard.
    anchors = ["E7", "E22", "A22"]
    for idx, spec in enumerate(chart_specs[:3]):
        kind, col, last_row, title = spec
        if kind == "line":
            chart = LineChart()
        else:
            chart = BarChart()
        chart.title = title[:45]
        chart.height = 7
        chart.width = 11
        data = Reference(ws, min_col=col + 1, min_row=1, max_row=last_row)
        cats_ref = Reference(ws, min_col=col, min_row=2, max_row=last_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats_ref)
        ws.add_chart(chart, anchors[idx])

    # Hide helper columns.
    for c in range(helper_col, helper_col + 10):
        ws.column_dimensions[get_column_letter(c)].hidden = True
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 55
    for c in range(4, 9):
        ws.column_dimensions[get_column_letter(c)].width = 14


def excel_report(df, business, kpis, insights, quality, filename=None):
    """Create a polished workbook, including an actual Dashboard sheet."""
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        # Dashboard first so it opens as the executive view.
        dashboard = writer.book.create_sheet("Dashboard", 0)
        _make_dashboard(dashboard, df, kpis, quality or {}, insights)

        _write_df(writer, df, "Data")
        _write_df(writer, business, "Business Analysis")

        kpi_rows = []
        for item in kpis or []:
            kpi_rows.append({
                "KPI": item.get("name", ""),
                "Value": item.get("value", ""),
                "Display Value": item.get("display_value", ""),
                "Description": item.get("description", ""),
                "Source Columns": ", ".join(item.get("source_columns", []) or []),
                "Format": item.get("format", ""),
            })
        _write_df(writer, pd.DataFrame(kpi_rows), "KPIs")
        _write_df(writer, pd.DataFrame({"Insight": insights or []}), "Insights")

        quality_summary = pd.DataFrame([
            {"Metric": "Health Score", "Value": (quality or {}).get("health_score", "")},
            {"Metric": "Missing Cells", "Value": (quality or {}).get("missing_cells", "")},
            {"Metric": "Duplicate Rows", "Value": (quality or {}).get("duplicate_rows", "")},
            {"Metric": "Potential Outlier Cells", "Value": (quality or {}).get("outlier_cells", "")},
        ])
        _write_df(writer, quality_summary, "Data Quality")

        missing = (quality or {}).get("missing_table")
        if isinstance(missing, pd.DataFrame):
            _write_df(writer, missing, "Missing Values")
        outliers = (quality or {}).get("outlier_table")
        if isinstance(outliers, pd.DataFrame):
            _write_df(writer, outliers, "Outliers")

    # Final workbook styling.
    output.seek(0)
    wb = load_workbook(output)
    if "Dashboard" in wb.sheetnames:
        wb.active = wb.sheetnames.index("Dashboard")
    for ws in wb.worksheets:
        if ws.title != "Dashboard":
            _style_sheet(ws)
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0

    final = BytesIO()
    wb.save(final)
    final.seek(0)
    return final.getvalue()
